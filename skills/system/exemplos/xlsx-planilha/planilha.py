#!/usr/bin/env python3
"""Cria uma planilha com formulas e grafico. Uso: planilha.py [saida.xlsx]."""
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


def criar(saida: Path, itens: list[tuple[str, int, float]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    ws.append(["Produto", "Quantidade", "Preço (R$)", "Total (R$)"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
        c.alignment = Alignment(horizontal="center")
    for i, (nome, qtd, preco) in enumerate(itens, start=2):
        ws.append([nome, qtd, preco, f"=B{i}*C{i}"])
    fim = len(itens) + 1
    ws.append(["Total", f"=SUM(B2:B{fim})", None, f"=SUM(D2:D{fim})"])
    ws.cell(row=fim + 1, column=1).font = Font(bold=True)
    for linha in ws.iter_rows(min_row=2, max_row=fim + 1, min_col=3, max_col=4):
        for c in linha:
            c.number_format = '"R$" #,##0.00'
    ws.column_dimensions["A"].width = 22
    for col in "BCD":
        ws.column_dimensions[col].width = 15
    ws.freeze_panes = "A2"

    g = BarChart()
    g.title = "Total por produto"
    g.y_axis.title = "R$"
    g.add_data(Reference(ws, min_col=4, min_row=1, max_row=fim), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=1, min_row=2, max_row=fim))
    ws.add_chart(g, "F2")

    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    return saida


if __name__ == "__main__":
    destino = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(os.environ.get("OUT_DIR", "/workspace/out")) / "vendas.xlsx"
    print(criar(destino, [("Camiseta", 10, 49.9), ("Boné", 5, 35.0), ("Caneca", 8, 22.5)]))
