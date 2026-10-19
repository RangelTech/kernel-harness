#!/usr/bin/env python3
"""Previa dos arquivos anexados. Uso: ler_pacote.py [pasta] (padrao /workspace/inbox)."""
import csv
import sys
from pathlib import Path


def previa_pdf(p: Path, paginas: int = 2) -> str:
    from pypdf import PdfReader

    texto = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages[:paginas]).strip()
    return texto[:1000] or "(sem texto extraivel: PDF provavelmente escaneado)"


def previa_xlsx(p: Path, linhas: int = 5) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(p, read_only=True, data_only=True)
    partes = []
    for ws in wb.worksheets[:3]:
        partes.append(f"[aba {ws.title}]")
        for i, linha in enumerate(ws.iter_rows(values_only=True)):
            if i >= linhas:
                break
            partes.append(" | ".join("" if v is None else str(v) for v in linha))
    return "\n".join(partes)


def previa_csv(p: Path, linhas: int = 5) -> str:
    with open(p, newline="", encoding="utf-8", errors="replace") as f:
        return "\n".join(" | ".join(r) for _, r in zip(range(linhas), csv.reader(f)))


def previa(p: Path) -> str:
    ext = p.suffix.lower()
    if ext == ".pdf":
        return previa_pdf(p)
    if ext in (".xlsx", ".xlsm"):
        return previa_xlsx(p)
    if ext == ".csv":
        return previa_csv(p)
    if ext in (".txt", ".md", ".json"):
        return p.read_text(encoding="utf-8", errors="replace")[:1000]
    return f"(binario {p.stat().st_size} bytes)"


def main(pasta: str) -> None:
    base = Path(pasta)
    if not base.is_dir():
        raise SystemExit(f"pasta inexistente: {base}. Use rt_env_upload para trazer o anexo.")
    for p in sorted(x for x in base.rglob("*") if x.is_file()):
        print(f"=== {p} ({p.stat().st_size} bytes)")
        print(previa(p))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "/workspace/inbox")
